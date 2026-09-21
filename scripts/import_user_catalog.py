#!/usr/bin/env python
"""scripts/import_user_catalog.py — Impor katalog klien dari XLSX (bertahap/chunked).

Dipakai untuk file besar (>5.000 baris) supaya tidak menabrak timeout HTTP: baris
dikelompokkan per PRODUK lalu dikirim ke `POST /api/admin/products/io/commit` dalam
batch kecil. Semua penyesuaian nilai (harga placeholder, konsentrasi, status) dilakukan
di sini agar file asli klien tidak diubah.

Pakai:
  python scripts/import_user_catalog.py <file.xlsx> [--price 1000] [--stock 0]
      [--concentration EDP] [--status archived] [--mode add-only] [--batch 100]
      [--limit N] [--dry-run]
"""
import argparse
import os
import sys
import time
from collections import OrderedDict

import requests
from openpyxl import load_workbook

BASE = os.environ.get("BACKEND_BASE", "http://localhost:8001/api")
ADMIN_EMAIL = os.environ.get("ADMIN_EMAIL", "admin@collectorparfum.id")
ADMIN_PASS = os.environ.get("ADMIN_PASS", "Admin#2026")

G, R, Y, C, X = "\033[92m", "\033[91m", "\033[93m", "\033[96m", "\033[0m"


def read_sheet(path, sheet="Produk"):
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet] if sheet in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    out = []
    for r in rows:
        if not any(c not in (None, "") for c in r):
            continue
        out.append({header[i]: ("" if v is None else str(v).strip()) for i, v in enumerate(r) if i < len(header)})
    wb.close()
    return header, out


def normalize(rows, price, stock, concentration, status):
    """Terapkan keputusan klien: harga placeholder, stok, konsentrasi, status."""
    for r in rows:
        if not str(r.get("variant_price") or "").strip() or float(r.get("variant_price") or 0) <= 0:
            r["variant_price"] = str(price)
        if not str(r.get("variant_stock") or "").strip():
            r["variant_stock"] = str(stock)
        else:
            r["variant_stock"] = str(int(float(r["variant_stock"])))
        if concentration and not str(r.get("concentration") or "").strip():
            r["concentration"] = concentration
        if status:
            r["status"] = status
    return rows


def group_by_product(rows):
    groups = OrderedDict()
    for r in rows:
        key = (r.get("slug") or "").strip() or (r.get("name") or "").strip()
        groups.setdefault(key, []).append(r)
    return groups


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("path")
    ap.add_argument("--price", type=int, default=1000)
    ap.add_argument("--stock", type=int, default=0)
    ap.add_argument("--concentration", default="EDP")
    ap.add_argument("--status", default="archived")
    ap.add_argument("--mode", default="add-only", choices=["add-only", "upsert"])
    ap.add_argument("--batch", type=int, default=100, help="jumlah PRODUK per request")
    ap.add_argument("--limit", type=int, default=0, help="batasi jumlah produk (uji coba)")
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()

    header, rows = read_sheet(a.path)
    rows = normalize(rows, a.price, a.stock, a.concentration, a.status)
    groups = group_by_product(rows)
    if a.limit:
        groups = OrderedDict(list(groups.items())[: a.limit])

    print(f"{C}File   :{X} {a.path}")
    print(f"{C}Kolom  :{X} {len(header)}")
    print(f"{C}Baris  :{X} {sum(len(v) for v in groups.values())}  ({len(groups)} produk)")
    print(f"{C}Setelan:{X} harga={a.price} stok={a.stock} konsentrasi={a.concentration} "
          f"status={a.status} mode={a.mode} batch={a.batch} produk/request")
    if a.dry_run:
        first = next(iter(groups.values()))
        print(f"{Y}DRY RUN — contoh baris pertama:{X}")
        for k in ("slug", "name", "category", "gender", "concentration", "status",
                  "option1_name", "option1_value", "option2_name", "option2_value",
                  "variant_price", "variant_stock"):
            print(f"   {k:26s} = {first[0].get(k)!r}")
        return 0

    tok = requests.post(f"{BASE}/auth/login",
                        json={"email": ADMIN_EMAIL, "password": ADMIN_PASS}, timeout=30)
    tok.raise_for_status()
    H = {"Authorization": f"Bearer {tok.json()['token']}"}
    mapping = {h: h for h in header if h}

    keys = list(groups)
    total = {"created": 0, "updated": 0, "skipped": 0, "failed": 0}
    errors = []
    t0 = time.time()
    for i in range(0, len(keys), a.batch):
        chunk_keys = keys[i:i + a.batch]
        chunk_rows = [r for k in chunk_keys for r in groups[k]]
        try:
            res = requests.post(f"{BASE}/admin/products/io/commit", headers=H, timeout=600,
                                json={"rows": chunk_rows, "mapping": mapping, "mode": a.mode})
        except Exception as e:
            print(f"  {R}batch {i//a.batch+1} GAGAL kirim: {e}{X}")
            errors.append(str(e))
            continue
        if res.status_code != 200:
            print(f"  {R}batch {i//a.batch+1} HTTP {res.status_code}: {res.text[:200]}{X}")
            errors.append(res.text[:200])
            continue
        d = res.json()
        for k in total:
            total[k] += int(d.get(k) or 0)
        errors.extend(d.get("errors") or [])
        done = min(i + a.batch, len(keys))
        print(f"  batch {i//a.batch+1:>3}/{(len(keys)+a.batch-1)//a.batch}  "
              f"produk {done}/{len(keys)}  +{d.get('created')} dibuat, "
              f"{d.get('updated')} diperbarui, {d.get('skipped')} dilewati, "
              f"{d.get('failed')} gagal   ({time.time()-t0:.0f}s)")

    print(f"\n{C}RINGKASAN{X}  dibuat={G}{total['created']}{X} diperbarui={total['updated']} "
          f"dilewati={total['skipped']} gagal={R if total['failed'] else ''}{total['failed']}{X} "
          f"— {time.time()-t0:.0f}s")
    if errors:
        print(f"{Y}Contoh error (maks 8):{X}")
        for e in errors[:8]:
            print("  -", e)
    return 0 if total["failed"] == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
